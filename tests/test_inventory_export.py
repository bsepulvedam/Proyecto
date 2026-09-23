import os
import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database.base import Base
from app.database.session import get_db
from app.main import app
from app.models.empresa import Empresa
from app.models.identity import Rol
from app.models.producto import Producto
from app.models.unidad_medida import UnidadMedida
from app.schemas.identity import UserCreate
from app.schemas.movimiento_inventario import LineaRecepcionCreate, RecepcionCreate
from app.services.auth_service import create_user
from app.services.inventario_movimiento_service import create_receipt
from tests.test_identity_auth import ASGIClient


class InventoryExportTests(unittest.TestCase):
    def setUp(self) -> None:
        names = ("APP_ENV", "AUTH_ENFORCED", "SESSION_SECRET", "COOKIE_SECURE")
        self.previous = {name: os.environ.get(name) for name in names}
        os.environ.update({
            "APP_ENV": "test",
            "AUTH_ENFORCED": "true",
            "SESSION_SECRET": "test-secret-only-not-production",
            "COOKIE_SECURE": "false",
        })
        self.engine = create_engine(
            "sqlite+pysqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine, expire_on_commit=False)
        with self.Session() as db:
            db.add(Rol(codigo="ADMIN", nombre="Administrador"))
            bol = Empresa(codigo="BOLIKLOR", nombre="BOLIKLOR")
            alm = Empresa(codigo="ALM", nombre="ALM")
            tineta = UnidadMedida(codigo="TINETA", nombre="Tineta", permite_decimales=False)
            db.add_all([bol, alm, tineta])
            db.commit()
            create_user(db, UserCreate(username="admin-export", password="Clave-Admin-Export-123"), "ADMIN")
            bol1 = Producto(
                empresa_id=bol.id, sku="BOL-1", nombre="PINTURA ACRILICA",
                unidad_stock_id=tineta.id, factor_conversion=1, unidad_costo_id=tineta.id, stock_minimo=10,
            )
            alm1 = Producto(
                empresa_id=alm.id, sku="ALM-1", nombre="PRODUCTO ALM",
                unidad_stock_id=tineta.id, factor_conversion=1, unidad_costo_id=tineta.id, stock_minimo=1,
            )
            db.add_all([bol1, alm1])
            db.commit()
            create_receipt(db, RecepcionCreate(
                empresa_id=bol.id, fecha=date(2026, 9, 1),
                lineas=[LineaRecepcionCreate(producto_id=bol1.id, cantidad_presentaciones=Decimal("5"), costo_unitario=Decimal("100"))],
            ))
            create_receipt(db, RecepcionCreate(
                empresa_id=alm.id, fecha=date(2026, 9, 1),
                lineas=[LineaRecepcionCreate(producto_id=alm1.id, cantidad_presentaciones=Decimal("3"), costo_unitario=Decimal("50"))],
            ))
            self.bol_id, self.alm_id = bol.id, alm.id

        def override_db():
            with self.Session() as db:
                yield db

        app.dependency_overrides[get_db] = override_db

    def tearDown(self) -> None:
        app.dependency_overrides.clear()
        self.engine.dispose()
        for name, value in self.previous.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value

    def admin_client(self):
        client = ASGIClient()
        page = client.get("/login")
        import re
        token = re.search(r'name="csrf_token" value="([^"]+)"', page.text).group(1)
        login = client.post("/login", {"username": "admin-export", "password": "Clave-Admin-Export-123", "csrf_token": token})
        self.assertEqual(login.status_code, 303)
        return client

    def test_export_without_authentication_redirects_to_login(self):
        response = ASGIClient().get("/inventario/movimientos/exportar")
        self.assertEqual(response.status_code, 303)

    def test_export_with_authentication_and_no_filters_returns_downloadable_xlsx(self):
        response = self.admin_client().get("/inventario/movimientos/exportar")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertRegex(
            response.headers["content-disposition"],
            r'^attachment; filename="movimientos_inventario_\d{8}_\d{6}\.xlsx"$',
        )
        self.assertTrue(len(response.body) > 0)
        workbook = load_workbook(BytesIO(response.body), read_only=True, data_only=False)
        sheet = workbook["Movimientos"]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][:6], ("Fecha", "N° documento", "Tipo", "Empresa", "SKU", "Producto"))
        self.assertEqual(len(rows) - 1, 2)
        skus = {row[4] for row in rows[1:]}
        self.assertEqual(skus, {"BOL-1", "ALM-1"})

    def test_export_filtered_by_company_only_includes_its_rows(self):
        response = self.admin_client().get(f"/inventario/movimientos/exportar?empresa_id={self.bol_id}")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.body), read_only=True)
        rows = list(workbook["Movimientos"].iter_rows(values_only=True))
        self.assertEqual(len(rows) - 1, 1)
        self.assertEqual(rows[1][3], "BOLIKLOR")
        self.assertEqual(rows[1][4], "BOL-1")


if __name__ == "__main__":
    unittest.main()
