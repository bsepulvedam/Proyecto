"""Carga inicial / reconciliación del inventario real exportado por el bot.

Lee el Excel real (`Maestro de Productos`, `Stock Consolidado`,
`Registro de Movimientos`) y lo concilia contra PostgreSQL: crea productos
nuevos sin sobreescribir los existentes, carga los movimientos históricos con
idempotencia por `ID Transacción`, y expone por separado la corrección de
saldo de BOL-12 (ver `apply_bol12_correction`).

No es un endpoint web: se invoca desde un script de management o desde tests.
Sigue el mismo patrón de análisis/ejecución que `product_import_service.py`,
pero es un módulo independiente — el formato de origen (export del bot) no
tiene relación con el Excel legacy que consume ese otro servicio.
"""
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models.bodega import Bodega
from app.models.empresa import Empresa
from app.models.movimiento_inventario import DetalleMovimientoInventario, MovimientoInventario
from app.models.producto import Producto
from app.models.unidad_medida import UnidadMedida
from app.services.inventario_movimiento_service import calculate_receipt_cost

logger = logging.getLogger(__name__)

SOURCE_FILENAME = "Inventario IA - PRUEBAS.xlsx"
SHEET_MAESTRO = "Maestro de Productos"
SHEET_STOCK = "Stock Consolidado"
SHEET_MOVIMIENTOS = "Registro de Movimientos"

UNIT_ALIASES = {"UNIDAD": "UN"}
BODEGA_TO_EMPRESA = {
    "BOLIKLOR": "BOLIKLOR",
    "ALM": "ALM",
    "MASV": "MASV",
    "MAS VIAL": "MASV",
}
TIPO_TO_LEDGER = {
    "RECEPCION": "RECEPCION",
    "DESPACHO": "DESPACHO",
    "DEVOLUCION": "DEVOLUCION",
}
ORIGEN_ERP_WEB = "ERP_WEB"
ORIGEN_BOT_TELEGRAM = "BOT_TELEGRAM"
ACTOR_MIGRACION_HISTORICA = "Migración histórica 2026-09-21"
BODEGA_PRINCIPAL_CODIGO = "PRINCIPAL"


class InventorySyncError(RuntimeError):
    """Errores de lectura/validación del Excel; no dejan datos a medio escribir."""


class InventorySyncExecutionError(RuntimeError):
    """Fallo técnico durante la escritura en PostgreSQL; la transacción se revierte."""


def default_source_path() -> Path:
    return Path("data/inventario_sync") / SOURCE_FILENAME


def _strip_accents(value: str) -> str:
    return unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).upper()


def normalize_key(value: object) -> str:
    """Normaliza para comparar contra los diccionarios de mapeo (sin tildes)."""
    return _strip_accents(normalize_text(value))


def normalize_unit(value: object) -> str:
    unit = normalize_text(value)
    return UNIT_ALIASES.get(unit, unit)


def parse_decimal(value: object) -> Decimal | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text == "-":
        return None
    try:
        return Decimal(text.replace(",", "."))
    except InvalidOperation:
        return None


def parse_fecha(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def company_from_sku(sku: str) -> str | None:
    for prefix, empresa_codigo in (("BOL-", "BOLIKLOR"), ("ALM-", "ALM"), ("MASV-", "MASV")):
        if sku.startswith(prefix):
            return empresa_codigo
    return None


@dataclass
class CatalogSyncRow:
    fila_excel: int
    sku: str
    empresa_code: str | None
    nombre: str
    unidad_stock: str | None
    familia: str
    stock_minimo: Decimal | None
    status: str = "NUEVO"  # NUEVO | EXISTENTE | CONFLICTO | ERROR
    issues: list[str] = field(default_factory=list)


@dataclass
class MovementSyncRow:
    fila_excel: int
    tx_id: str
    fecha: date | None
    tipo_excel: str
    tipo_ledger: str | None
    bodega_excel: str
    empresa_code: str | None
    sku: str
    cantidad: Decimal | None
    costo_unitario: Decimal | None
    localidad: str | None
    observacion: str | None
    usuario_origen: str
    origen: str
    actor_referencia: str
    status: str = "VALIDO"  # VALIDO | ERROR
    issues: list[str] = field(default_factory=list)


@dataclass
class InventorySyncReport:
    source_file: str
    catalog_rows: list[CatalogSyncRow]
    movement_rows: list[MovementSyncRow]
    stock_sin_maestro: list[str]

    @property
    def catalog_nuevos(self) -> list[CatalogSyncRow]:
        return [row for row in self.catalog_rows if row.status == "NUEVO"]

    @property
    def catalog_existentes(self) -> list[CatalogSyncRow]:
        return [row for row in self.catalog_rows if row.status == "EXISTENTE"]

    @property
    def catalog_conflictos(self) -> list[CatalogSyncRow]:
        return [row for row in self.catalog_rows if row.status == "CONFLICTO"]

    @property
    def catalog_errores(self) -> list[CatalogSyncRow]:
        return [row for row in self.catalog_rows if row.status == "ERROR"]

    @property
    def movimientos_validos(self) -> list[MovementSyncRow]:
        return [row for row in self.movement_rows if row.status == "VALIDO"]

    @property
    def movimientos_con_error(self) -> list[MovementSyncRow]:
        return [row for row in self.movement_rows if row.status == "ERROR"]


@dataclass
class CatalogSyncResult:
    creados: int
    ya_existian: int
    fallidos: list[str]


@dataclass
class MovementSyncResult:
    cargados: int
    ya_existian: int
    fallidos: list[str]


def _find_sheet(workbook, name: str):
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().casefold() == name.casefold():
            return workbook[sheet_name]
    raise InventorySyncError(f'No se encontró la hoja requerida "{name}".')


def _existing_product_conflicts(row: CatalogSyncRow, existing: Producto) -> list[str]:
    conflicts: list[str] = []
    if existing.empresa.codigo != row.empresa_code:
        conflicts.append("La empresa difiere del producto ya existente en Postgres.")
    if normalize_text(existing.nombre) != row.nombre:
        conflicts.append("El nombre difiere del producto ya existente en Postgres.")
    if existing.unidad_stock.codigo != row.unidad_stock:
        conflicts.append("La unidad de stock difiere del producto ya existente en Postgres.")
    if row.stock_minimo is not None and existing.stock_minimo != row.stock_minimo:
        conflicts.append("El stock mínimo difiere del producto ya existente en Postgres.")
    if row.familia and normalize_text(existing.familia) != row.familia:
        conflicts.append("La familia difiere del producto ya existente en Postgres.")
    return conflicts


def analyze_inventory_workbook(db: Session, source_path: Path | None = None) -> InventorySyncReport:
    """Lee y valida el Excel real del bot sin escribir en la base de datos.

    Separa el análisis de la ejecución (mismo patrón que
    ``product_import_service.analyze_product_workbook``) para poder inspeccionar
    conflictos antes de decidir si se ejecuta ``sync_catalog``/``sync_movements``.
    """
    path = source_path or default_source_path()
    if not path.is_file():
        raise InventorySyncError(f"No se encontró el archivo fuente {path}.")
    try:
        workbook = load_workbook(path, data_only=True)
    except (OSError, InvalidFileException, BadZipFile) as exc:
        raise InventorySyncError("El archivo fuente no es un libro Excel válido.") from exc

    try:
        maestro = _find_sheet(workbook, SHEET_MAESTRO)
        stock = _find_sheet(workbook, SHEET_STOCK)
        movimientos = _find_sheet(workbook, SHEET_MOVIMIENTOS)

        known_units = {unit.codigo for unit in db.scalars(select(UnidadMedida)).all()}
        known_companies = {empresa.codigo for empresa in db.scalars(select(Empresa)).all()}
        existing_products = {
            product.sku: product
            for product in db.scalars(
                select(Producto).options(
                    joinedload(Producto.empresa),
                    joinedload(Producto.unidad_stock),
                )
            ).all()
        }

        stock_minimo_by_sku: dict[str, Decimal] = {}
        stock_skus: set[str] = set()
        for values in stock.iter_rows(min_row=2, values_only=True):
            if not values or not values[0]:
                continue
            sku = normalize_text(values[0])
            stock_skus.add(sku)
            minimo = parse_decimal(values[3])
            if minimo is not None:
                stock_minimo_by_sku[sku] = minimo

        catalog_rows: list[CatalogSyncRow] = []
        maestro_skus: set[str] = set()
        for excel_row, values in enumerate(maestro.iter_rows(min_row=2, values_only=True), start=2):
            if not values or not values[0]:
                continue
            sku = normalize_text(values[0])
            maestro_skus.add(sku)
            row = CatalogSyncRow(
                fila_excel=excel_row,
                sku=sku,
                empresa_code=company_from_sku(sku),
                nombre=normalize_text(values[1]),
                unidad_stock=normalize_unit(values[2]) or None,
                familia=normalize_text(values[3]),
                stock_minimo=stock_minimo_by_sku.get(sku),
            )
            existing = existing_products.get(sku)
            if existing is not None:
                conflicts = _existing_product_conflicts(row, existing)
                row.status = "CONFLICTO" if conflicts else "EXISTENTE"
                row.issues.extend(conflicts)
            else:
                if row.empresa_code is None:
                    row.issues.append("No se pudo identificar la empresa a partir del SKU.")
                elif row.empresa_code not in known_companies:
                    row.issues.append(f"Empresa {row.empresa_code} no registrada en Postgres.")
                if not row.unidad_stock or row.unidad_stock not in known_units:
                    row.issues.append(f"Unidad no reconocida: {values[2]!r}.")
                if not row.nombre:
                    row.issues.append("Nombre vacío.")
                row.status = "ERROR" if row.issues else "NUEVO"
            catalog_rows.append(row)

        movement_rows: list[MovementSyncRow] = []
        for excel_row, values in enumerate(movimientos.iter_rows(min_row=2, values_only=True), start=2):
            if not values or not values[0]:
                continue
            tx_id = str(values[0]).strip()
            fecha = parse_fecha(values[1])
            tipo_excel = normalize_text(values[2])
            tipo_ledger = TIPO_TO_LEDGER.get(normalize_key(tipo_excel))
            bodega_excel = normalize_text(values[3])
            empresa_code = BODEGA_TO_EMPRESA.get(normalize_key(bodega_excel))
            sku = normalize_text(values[4])
            cantidad = parse_decimal(values[6])
            costo_unitario = parse_decimal(values[7])
            localidad = normalize_text(values[8]) or None
            observacion = normalize_text(values[9]) or None
            usuario_origen = normalize_text(values[10])
            origen = ORIGEN_BOT_TELEGRAM if usuario_origen == "BRUNO" else ORIGEN_ERP_WEB
            actor_referencia = (
                "Bruno (Telegram)" if origen == ORIGEN_BOT_TELEGRAM else ACTOR_MIGRACION_HISTORICA
            )

            row = MovementSyncRow(
                fila_excel=excel_row, tx_id=tx_id, fecha=fecha, tipo_excel=tipo_excel,
                tipo_ledger=tipo_ledger, bodega_excel=bodega_excel, empresa_code=empresa_code,
                sku=sku, cantidad=cantidad, costo_unitario=costo_unitario,
                localidad=localidad, observacion=observacion, usuario_origen=usuario_origen,
                origen=origen, actor_referencia=actor_referencia,
            )
            if not tx_id:
                row.issues.append("ID Transacción vacío.")
            elif len(tx_id) > 30:
                row.issues.append("ID Transacción excede 30 caracteres (límite de numero_documento).")
            if fecha is None:
                row.issues.append("Fecha vacía o no reconocible.")
            if tipo_ledger is None:
                row.issues.append(f"Tipo de movimiento no reconocido: {tipo_excel!r}.")
            if empresa_code is None:
                row.issues.append(f"Bodega no reconocida: {bodega_excel!r}.")
            elif empresa_code not in known_companies:
                row.issues.append(f"Empresa {empresa_code} no registrada en Postgres.")
            if not sku:
                row.issues.append("SKU vacío.")
            elif sku not in maestro_skus and sku not in existing_products:
                row.issues.append(f"SKU {sku} no existe en el catálogo (Maestro ni Postgres).")
            if cantidad is None or cantidad <= 0:
                row.issues.append("Cantidad vacía, no numérica o no positiva.")
            if row.issues:
                row.status = "ERROR"
            movement_rows.append(row)

        return InventorySyncReport(
            source_file=path.name,
            catalog_rows=catalog_rows,
            movement_rows=movement_rows,
            stock_sin_maestro=sorted(stock_skus - maestro_skus),
        )
    finally:
        workbook.close()


def sync_catalog(db: Session, report: InventorySyncReport) -> CatalogSyncResult:
    """Crea únicamente los productos NUEVO del reporte, en un solo commit.

    Nunca sobreescribe un producto ya existente: los CONFLICTO/ERROR quedan
    fuera y se reportan, no se resuelven silenciosamente.
    """
    nuevos = report.catalog_nuevos
    logger.info(
        "Inicio carga de catálogo: nuevos=%s existentes=%s conflictos=%s errores=%s",
        len(nuevos), len(report.catalog_existentes), len(report.catalog_conflictos), len(report.catalog_errores),
    )
    try:
        companies = {empresa.codigo: empresa for empresa in db.scalars(select(Empresa)).all()}
        units = {unit.codigo: unit for unit in db.scalars(select(UnidadMedida)).all()}
        existing_skus = {sku for (sku,) in db.execute(select(Producto.sku))}
        creados = 0
        fallidos: list[str] = []
        for row in nuevos:
            if row.sku in existing_skus:
                continue  # ya se creó en una corrida anterior; idempotente
            empresa = companies.get(row.empresa_code or "")
            unidad = units.get(row.unidad_stock or "")
            if empresa is None or unidad is None:
                fallidos.append(f"{row.sku}: empresa o unidad ya no resoluble al momento de crear.")
                continue
            db.add(Producto(
                empresa_id=empresa.id, sku=row.sku, nombre=row.nombre,
                unidad_stock_id=unidad.id, unidad_contenido_id=None,
                factor_conversion=Decimal("1"), unidad_costo_id=unidad.id,
                stock_minimo=row.stock_minimo if row.stock_minimo is not None else Decimal("0"),
                tipo=None, familia=row.familia or None, activo=True,
            ))
            existing_skus.add(row.sku)
            creados += 1
        db.commit()
        result = CatalogSyncResult(creados=creados, ya_existian=len(report.catalog_existentes), fallidos=fallidos)
        logger.info("Fin carga de catálogo: creados=%s fallidos=%s", result.creados, len(result.fallidos))
        return result
    except Exception as exc:
        db.rollback()
        logger.exception("Error técnico durante la carga de catálogo")
        raise InventorySyncExecutionError("No fue posible completar la carga de catálogo.") from exc


def sync_movements(db: Session, report: InventorySyncReport) -> MovementSyncResult:
    """Carga los movimientos VALIDO del reporte, idempotente por ``numero_documento``.

    Requiere haber corrido ``sync_catalog`` primero si el reporte incluye
    productos NUEVO referenciados por movimientos.
    """
    validos = report.movimientos_validos
    logger.info(
        "Inicio carga de movimientos: validos=%s con_error=%s", len(validos), len(report.movimientos_con_error),
    )
    try:
        products = {
            product.sku: product
            for product in db.scalars(
                select(Producto).options(
                    joinedload(Producto.unidad_stock),
                    joinedload(Producto.unidad_contenido),
                    joinedload(Producto.unidad_costo),
                )
            ).all()
        }
        companies = {empresa.codigo: empresa for empresa in db.scalars(select(Empresa)).all()}
        bodegas = {
            (bodega.empresa_id, bodega.codigo): bodega for bodega in db.scalars(select(Bodega)).all()
        }
        ya_cargados = {numero for (numero,) in db.execute(select(MovimientoInventario.numero_documento))}

        cargados = 0
        ya_existian = 0
        fallidos: list[str] = []
        for row in validos:
            if row.tx_id in ya_cargados:
                ya_existian += 1
                continue
            empresa = companies.get(row.empresa_code)
            product = products.get(row.sku)
            bodega = bodegas.get((empresa.id, BODEGA_PRINCIPAL_CODIGO)) if empresa else None
            if empresa is None or product is None or bodega is None:
                fallidos.append(f"{row.tx_id}: empresa, producto o bodega ya no resoluble al momento de cargar.")
                continue
            if product.empresa_id != empresa.id:
                fallidos.append(f"{row.tx_id}: {row.sku} pertenece a una empresa distinta de {row.empresa_code}.")
                continue

            costo_unitario = row.costo_unitario if row.tipo_ledger == "RECEPCION" else None
            if costo_unitario is not None:
                presentation_cost, total = calculate_receipt_cost(product, row.cantidad, costo_unitario)
            else:
                presentation_cost = total = None

            movement = MovimientoInventario(
                tipo=row.tipo_ledger, empresa_id=empresa.id, bodega_id=bodega.id,
                fecha=row.fecha, numero_documento=row.tx_id,
                comuna=row.localidad, observaciones=row.observacion,
                origen=row.origen, actor_referencia=row.actor_referencia,
            )
            movement.detalles.append(DetalleMovimientoInventario(
                producto_id=product.id, cantidad_presentaciones=row.cantidad,
                unidad_presentacion_snapshot=product.unidad_stock.codigo,
                factor_conversion_snapshot=product.factor_conversion,
                unidad_contenido_snapshot=product.unidad_contenido.codigo if product.unidad_contenido else None,
                unidad_costo_snapshot=product.unidad_costo.codigo if product.unidad_costo else None,
                costo_unitario=costo_unitario, costo_presentacion=presentation_cost, valor_total=total,
            ))
            db.add(movement)
            ya_cargados.add(row.tx_id)
            cargados += 1
        db.commit()
        result = MovementSyncResult(cargados=cargados, ya_existian=ya_existian, fallidos=fallidos)
        logger.info("Fin carga de movimientos: cargados=%s ya_existian=%s fallidos=%s",
                     result.cargados, result.ya_existian, len(result.fallidos))
        return result
    except Exception as exc:
        db.rollback()
        logger.exception("Error técnico durante la carga de movimientos")
        raise InventorySyncExecutionError("No fue posible completar la carga de movimientos.") from exc


def apply_bol12_correction(
    db: Session,
    adjustment_quantity: Decimal,
    numero_documento: str = "AJU-BOL12-2026-09-21",
) -> MovimientoInventario:
    """Aplica la corrección de saldo de BOL-12 (Boliklor) como AJUSTE_POSITIVO.

    El material sí llegó a Boliklor: el saldo negativo es un error de registro
    durante las pruebas del bot, no un préstamo a otra empresa (confirmado por
    el usuario, 2026-09-21) — por eso es ``AJUSTE_POSITIVO`` y no
    ``TRANSFERENCIA_ENTRE_EMPRESAS``.

    ``adjustment_quantity`` es obligatorio y sin default de negocio: el saldo
    neto de los 278 movimientos de ``Registro de Movimientos`` para BOL-12 da
    -60, no -40 (el -40 de ``Stock Consolidado`` incluye una operación de
    prueba del bot que nunca quedó registrada como movimiento formal — ver
    hallazgo reportado 2026-09-21). El monto exacto debe confirmarse antes de
    invocar esta función; no se infiere aquí.
    """
    if adjustment_quantity <= 0:
        raise InventorySyncError("El ajuste de corrección debe ser una cantidad positiva.")

    existing = db.scalar(
        select(MovimientoInventario).where(MovimientoInventario.numero_documento == numero_documento)
    )
    if existing is not None:
        return existing  # idempotente: ya se aplicó en una corrida anterior

    empresa = db.scalar(select(Empresa).where(Empresa.codigo == "BOLIKLOR"))
    product = db.scalar(
        select(Producto).options(joinedload(Producto.unidad_stock)).where(Producto.sku == "BOL-12")
    )
    bodega = (
        db.scalar(
            select(Bodega).where(Bodega.empresa_id == empresa.id, Bodega.codigo == BODEGA_PRINCIPAL_CODIGO)
        )
        if empresa is not None
        else None
    )
    if empresa is None or product is None or bodega is None:
        raise InventorySyncError("No fue posible resolver empresa/bodega/producto BOL-12 para la corrección.")

    motivo = "Corrección de saldo: error de registro durante pruebas del bot (2026-09-21)"
    try:
        movement = MovimientoInventario(
            tipo="AJUSTE_POSITIVO", empresa_id=empresa.id, bodega_id=bodega.id,
            fecha=date(2026, 9, 21), numero_documento=numero_documento,
            observaciones=motivo, origen=ORIGEN_ERP_WEB, actor_referencia=motivo,
        )
        movement.detalles.append(DetalleMovimientoInventario(
            producto_id=product.id, cantidad_presentaciones=adjustment_quantity,
            unidad_presentacion_snapshot=product.unidad_stock.codigo,
            factor_conversion_snapshot=product.factor_conversion,
        ))
        db.add(movement)
        db.commit()
        return movement
    except Exception as exc:
        db.rollback()
        logger.exception("Error técnico aplicando la corrección de BOL-12")
        raise InventorySyncExecutionError("No fue posible aplicar la corrección de BOL-12.") from exc


def print_sync_summary(
    report: InventorySyncReport,
    catalog_result: CatalogSyncResult,
    movement_result: MovementSyncResult,
) -> None:
    """Reporte legible por consola: qué se cargó, qué ya existía, qué conflictos hubo."""
    print(f"Archivo: {report.source_file}")
    print(
        f"Catálogo — creados: {catalog_result.creados} | ya existían: {catalog_result.ya_existian} | "
        f"conflictos: {len(report.catalog_conflictos)} | errores: {len(report.catalog_errores) + len(catalog_result.fallidos)}"
    )
    for row in report.catalog_conflictos:
        print(f"  CONFLICTO {row.sku}: {'; '.join(row.issues)}")
    for row in report.catalog_errores:
        print(f"  ERROR {row.sku}: {'; '.join(row.issues)}")
    for message in catalog_result.fallidos:
        print(f"  ERROR {message}")
    if report.stock_sin_maestro:
        print(f"  SKUs en Stock Consolidado sin fila en Maestro (no cargados): {', '.join(report.stock_sin_maestro)}")

    print(
        f"Movimientos — cargados: {movement_result.cargados} | ya existían: {movement_result.ya_existian} | "
        f"rechazados: {len(report.movimientos_con_error) + len(movement_result.fallidos)}"
    )
    for row in report.movimientos_con_error:
        print(f"  ERROR {row.tx_id}: {'; '.join(row.issues)}")
    for message in movement_result.fallidos:
        print(f"  ERROR {message}")


def run_initial_sync(
    db: Session, source_path: Path | None = None
) -> tuple[InventorySyncReport, CatalogSyncResult, MovementSyncResult]:
    """Punto de entrada único: analiza, carga catálogo y luego movimientos.

    No aplica la corrección de BOL-12 — requiere una cantidad confirmada
    explícitamente (ver ``apply_bol12_correction``), así que se invoca aparte.
    """
    report = analyze_inventory_workbook(db, source_path)
    catalog_result = sync_catalog(db, report)
    movement_result = sync_movements(db, report)
    return report, catalog_result, movement_result
